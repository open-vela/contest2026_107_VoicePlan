from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "VelaPlan-作品介绍文档.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size=11, color="1F2937"):
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)


def set_spacing(style, before=0, after=6, line=1.1):
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(paragraph, fill="F4F6F9"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.add_run(text)
    return p


def add_bullet(doc, text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_status_box(doc, title, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    set_cell_shading(p)
    run = p.add_run(title + "：")
    set_run_font(run, bold=True, color="1F4D78")
    run = p.add_run(text)
    set_run_font(run, color="374151")
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
set_style_font(normal, size=10.5)
set_spacing(normal, after=4, line=1.05)

for name, size, color, before, after in [
    ("Heading 1", 16, "000000", 12, 6),
    ("Heading 2", 13, "000000", 8, 4),
    ("Heading 3", 12, "000000", 6, 3),
]:
    style = doc.styles[name]
    set_style_font(style, size=size, color=color)
    style.font.bold = True
    set_spacing(style, before=before, after=after, line=1.1)
    style.paragraph_format.keep_with_next = True

for list_name in ("List Bullet", "List Number"):
    style = doc.styles[list_name]
    set_style_font(style, size=10.5)
    set_spacing(style, after=2, line=1.08)
    style.paragraph_format.left_indent = Inches(0.5)
    style.paragraph_format.first_line_indent = Inches(-0.25)

title_style = doc.styles["Title"]
set_style_font(title_style, size=26, color="000000")
title_style.font.bold = True
title_style_ppr = title_style._element.get_or_add_pPr()
title_border = title_style_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_style_ppr.remove(title_border)
title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(4)
run = title.add_run("VelaPlan 腕上 AI 多场景计划助手")
set_run_font(run, size=26, bold=True, color="000000")

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.add_run("2026 首届 openvela AI 硬件开发者大赛 | 作品介绍文档")
set_run_font(run, size=12, color="6B7280")

add_heading(doc, "版本与验证结论", 2)
add_body(doc, "截至 2026 年 9 月 18 日，快应用、健康数据、真实 MiMo 语音转写、本地规则计划和发布版 RPK 均已验证。6 组 JavaScript 测试、3 个日志导出器测试及官方日志校验通过。system.velaclaw 仍有工具链警告，演示必须保留本地规则兜底提示；当前只缺 Demo 视频。")

add_heading(doc, "一、项目基本信息", 1)
add_heading(doc, "项目名称", 2)
add_body(doc, "VelaPlan 腕上 AI 多场景计划助手")
add_heading(doc, "参赛方向", 2)
add_body(doc, "手表应用创新。作品基于 openvela 快应用，面向智能手表/手环提供轻量化计划创建、腕上执行和状态感知调整。")
add_heading(doc, "代码仓库", 2)
add_body(doc, "https://github.com/open-vela/contest2026_107_VoicePlan")
add_body(doc, "源码基于 dev-ai-contest-2026 开发。当前开发分支为 voiceplan-work；官方仓库 PR #1 为 Open、可合并状态，CLA 检查已通过，仍需维护者审核并合入对应分支。")

add_heading(doc, "二、产品目标与应用场景", 1)
add_heading(doc, "目标用户", 2)
add_body(doc, "大学生、备赛团队成员、健身初学者，以及需要低打扰日程管理的日常用户。")
add_heading(doc, "解决的问题", 2)
add_body(doc, "普通待办工具主要记录任务，不能结合用户当天状态调整任务密度和运动强度；手机端创建和执行计划的操作链路也较长。VelaPlan 将“今天做什么”压缩到腕上交互，并把心率、血氧、压力转化为可解释的计划建议。")
add_heading(doc, "核心流程", 2)
add_body(doc, "输入一个统一目标并选择计划周期 → 可标记重要事项和提醒时间 → 通过文字确认或录制语音 → 结合健康和天气生成结构化计划 → 查看时间、时长、建议和调整原因 → 腕上完成、调整或延期任务 → 根据完成情况复盘并给出次日建议。")

add_heading(doc, "三、功能与创新点", 1)
for item in [
    "统一目标计划：不要求用户先判断任务属于日常、学习或健身，直接输入目标并按今日、本周、本月、本季度组织任务。",
    "状态感知计划：压力偏高时加入呼吸放松并降低任务密度；心率相对个人基线偏高时把运动调整为低强度活动；血氧偏低时安排恢复性活动；雨、大风和高温场景会调整户外任务。",
    "腕上执行闭环：任务按时间、类型、时长和建议展示，支持完成打卡，复盘区实时反馈完成数量和次日建议。",
    "AI 与规则双通道：优先通过 @system.velaclaw 调用设备端 ai_agent 生成结构化 JSON；AI 不可用或返回格式不符合约束时，使用本地规则计划保证核心流程可演示。",
    "语音入口：通过官方 @system.record 录制 WAV 音频，经 @system.file 读取并由 @system.fetch 调用 MiMo 多模态接口转写；文字进入输入框后由用户确认，再生成计划。",
]:
    add_bullet(doc, item)

add_heading(doc, "四、openvela 系统能力使用", 1)
for item in [
    "图形能力：使用 openvela 快应用页面、滚动容器、文本、输入框和交互控件，完成健康卡片、周期/提醒/天气设置、计划任务和复盘界面。",
    "AI 能力：使用 @system.velaclaw 的 ask 接口，将目标、周期、重要提醒、天气、心情、健康状态和结构化输出约束传给设备端 AI Agent。",
    "多媒体能力：使用 @system.record 的 start / stop 接口，采集单声道 16 kHz WAV 音频。",
    "健康能力：使用 @service.health 的 getRecentSamples、subscribeSample 和 unsubscribeSample，读取并订阅心率、血氧、压力数据。",
    "后台能力：在 src/manifest.json 中声明 config.background.features，使健康订阅具备后台持续回调条件。",
]:
    add_bullet(doc, item)
add_heading(doc, "技术优化与扩展建议", 2)
for item in [
    "用本地规则先生成可解释的最小计划，再用 AI 结果替换，降低网络和 AI 服务不稳定对演示的影响。",
    "要求 AI 输出固定 JSON，并在端侧验证 tasks、任务时间、名称和时长，避免自由文本造成腕上渲染失败。",
    "只在端侧展示非医疗性质的状态建议，不输出诊断或治疗结论。",
    "后续可增加任务逐项状态、历史完成率、ASR 纠错和圆形表盘适配；进入决赛后再接真实设备传感器。",
]:
    add_bullet(doc, item)

add_heading(doc, "五、技术方案与硬件关系", 1)
add_body(doc, "端侧使用 quickapp/hello_quickapp/ 快应用实现 UI、健康数据读取、录音入口、计划展示和复盘。planner.js 负责解析可用时间、判断状态和生成本地兜底计划；设备端 AI 负责生成同一输出结构的个性化计划。本地 backend/ 提供浏览器预览、MiMo Chat Completions 联调和服务失败时的规则兜底。")
add_body(doc, "初赛以官方 vela-miwear-watch-5.0（开发者大赛）模拟器和 service.health Mock 数据完成验证。SF32LB52-DevKit-LCD 与 1.85 英寸 390x450 AMOLED 屏幕用于展示屏幕、触摸、按键和腕上交互原型；本项目不把底层固件移植和真实传感器接入作为初赛主线。")

add_heading(doc, "六、AI-Native 开发说明", 1)
add_body(doc, "AI 参与了需求拆解、赛道选择、技术路线、快应用页面、计划规则、MiMo Prompt、测试、文档和调试问题定位。开发过程保留“先本地规则、再 AI 替换、失败可回退”的工程约束，并通过单元测试验证健康状态与计划输出。")
add_bullet(doc, "AI Coding 代码占比：约 90%。统计口径为核心快应用代码、测试、构建脚本和文档中由 AI 生成或修改、再由参赛者确认的内容；需求、账号操作、语音输入和视频录制由参赛者完成，因此该数值是过程估算，不等同于 Git 行数归属。")
add_bullet(doc, "使用的 AI 工具、MCP 和 Skills：Codex Desktop 为主开发工具，Claude Code 用于日志链路环境检查；本机命令与应用自动化工具用于 Git、AIoT-IDE、ADB、构建和验证。项目沉淀 Skill 为 docs/skills/health-aware-watch-planning/SKILL.md，并采用系统化调试和测试先行流程。")
add_bullet(doc, "Token 消耗：Codex 134,334,527，Claude Code 3,509，合计 134,338,036 tokens，原始统计见 logs/Yjwqj/manifest.json。MiMo 语音已真实成功调用，音频 Token 和费用以平台账单为准。")
add_bullet(doc, "完整 AI Coding 日志：logs/Yjwqj/ 包含 2 个真实会话，Codex 长会话由原始 rollout 脱敏转换为组委会 schema；组委会 validate-log.py 校验为 ALL OK，日志包保存为 submission/AI-Coding-logs.zip。")

add_heading(doc, "七、运行与验证", 1)
add_body(doc, "在仓库根目录执行：")
code = doc.add_paragraph()
code.paragraph_format.left_indent = Inches(0.25)
code.paragraph_format.space_after = Pt(8)
set_cell_shading(code, "F2F4F7")
run = code.add_run("npm test")
set_run_font(run, color="1F2937")
add_body(doc, "截至 2026 年 9 月 18 日，npm test 的 planner、plan store、MiMo 客户端、页面布局、语音转写和语音页面生命周期 6 组测试均通过；Codex 日志导出器的 3 个 Python 回归测试通过；npm run release 已生成生产 RPK。")
add_body(doc, "用 AIoT-IDE 打开 quickapp/hello_quickapp/，选择 VelaPlan_390x450 模拟器，点击调试并按“示例 → 重要事项/提醒 → 下雨 → 生成计划 → 完成一项”的顺序演示。提交包使用：")
code = doc.add_paragraph()
code.paragraph_format.left_indent = Inches(0.25)
code.paragraph_format.space_after = Pt(8)
set_cell_shading(code, "F2F4F7")
run = code.add_run("cd quickapp/hello_quickapp\nnpm run release")
set_run_font(run, color="1F2937")
add_body(doc, "发布包 submission/VelaPlan.release.rpk 的 SHA-256 为 DB5E3665DE346D6BE54EF0CF5CB8164F284BC0EC53143A0B14E8339EDC012082，设备端真实 MiMo 语音转写已验收。最终视频须展示语音目标、状态感知计划、执行和复盘，并控制在 5 分钟以内。")

add_heading(doc, "八、合规与提交前事项", 1)
add_body(doc, "作品遵循 Apache 2.0，不提交密钥、私钥、.env、编译缓存或未授权素材，不宣传医疗诊断。真实语音转写已经验收；当前只需录制 5 分钟内视频并等待 PR #1 合入。")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("VelaPlan | openvela AI 硬件开发者大赛 | ")
set_run_font(run, size=9, color="6B7280")
add_page_field(footer)

doc.core_properties.title = "VelaPlan 腕上 AI 多场景计划助手"
doc.core_properties.subject = "openvela 手表应用创新参赛作品介绍"
doc.core_properties.author = "VelaPlan Team"
doc.core_properties.keywords = "openvela, VelaPlan, AI, watch, quickapp"
doc.save(OUTPUT)
print(OUTPUT)
